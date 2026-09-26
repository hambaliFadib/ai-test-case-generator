"""Parser for Microsoft Word DOCX requirement documents."""

from pathlib import Path

from models.input_model import ParsedInput


class DocxParseError(ValueError):
    """Raised when a DOCX document cannot be parsed or is empty."""


def parse_docx(path: str | Path) -> ParsedInput:
    """Extract paragraphs and table rows from a DOCX file."""

    file_path = Path(path)
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError as exc:
        raise DocxParseError("python-docx is required for DOCX input.") from exc

    try:
        document = Document(str(file_path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        tables: list[str] = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    tables.append(" | ".join(cells))
    except (OSError, ValueError, KeyError, PackageNotFoundError) as exc:
        raise DocxParseError(f"Could not parse DOCX file '{file_path}': {exc}") from exc

    sections = paragraphs + tables
    text = "\n".join(sections).strip()
    if not text:
        raise DocxParseError(f"DOCX file '{file_path}' contains no readable text.")
    return ParsedInput(
        source_type="docx",
        source_name=str(file_path),
        text=text,
        metadata={"paragraph_count": len(paragraphs), "table_row_count": len(tables)},
    )
